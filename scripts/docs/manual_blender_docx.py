#!/usr/bin/env python
"""Gera o DOCX do Manual Blender para agentes a partir do Markdown (fonte única).

    python scripts/docs/manual_blender_docx.py            # docs/MANUAL_BLENDER_AGENTES.md -> docs/ARENAMON_Manual_Blender_Para_Agentes_IA_V2.docx

Usa o DOCX V1 original como MODELO DE ESTILO (fontes, cores, margens): o corpo é apagado e reescrito; o arquivo V1 não é
alterado. Converte títulos, parágrafos, listas, tabelas, blocos de código, citações e negrito/código em linha. Links viram
texto (o destino aparece entre parênteses quando é URL externa). Requer python-docx.
"""
import argparse
import os
import re
import sys

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
MD = os.path.join(RAIZ, "docs", "MANUAL_BLENDER_AGENTES.md")
MODELO = os.path.join(RAIZ, "docs", "ARENAMON_Manual_Blender_Para_Agentes_IA_V1_original.docx")
SAIDA = os.path.join(RAIZ, "docs", "ARENAMON_Manual_Blender_Para_Agentes_IA_V2.docx")

CABECALHO_TABELA = "18354A"   # mesma cor do V1
ALERTAS = {"NOTE": "Nota", "TIP": "Dica", "IMPORTANT": "Importante", "WARNING": "Atenção", "CAUTION": "Cuidado"}
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def limpar_corpo(doc):
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    for el in list(body):
        if el is not sect:
            body.remove(el)


def texto_em_linha(par, texto, negrito=False, cor=None, tamanho=None):
    """Escreve `texto` no parágrafo tratando **negrito**, `código` e [link](destino)."""
    for parte in INLINE.split(texto):
        if not parte:
            continue
        b, mono = negrito, False
        if parte.startswith("**") and parte.endswith("**"):
            parte, b = parte[2:-2], True
        elif parte.startswith("`") and parte.endswith("`"):
            parte, mono = parte[1:-1], True
        elif parte.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", parte)
            rotulo, destino = m.group(1), m.group(2)
            parte = rotulo if not destino.startswith("http") else "%s (%s)" % (rotulo, destino)
            parte = re.sub(r"`([^`]+)`", r"\1", parte)
        run = par.add_run(parte)
        run.bold = b or None
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        if cor:
            run.font.color.rgb = RGBColor.from_string(cor)
        if tamanho:
            run.font.size = tamanho


def sombrear(celula, cor):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcPr.append(shd)


def tabela(doc, linhas):
    celulas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas]
    celulas = [r for r in celulas if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    ncol = max(len(r) for r in celulas)
    t = doc.add_table(rows=len(celulas), cols=ncol)
    t.style = doc.styles["Table Grid"]
    for i, r in enumerate(celulas):
        for j in range(ncol):
            c = t.cell(i, j)
            p = c.paragraphs[0]
            if i == 0:
                sombrear(c, CABECALHO_TABELA)
                texto_em_linha(p, r[j] if j < len(r) else "", negrito=True, cor="FFFFFF", tamanho=Pt(8.5))
            else:
                texto_em_linha(p, r[j] if j < len(r) else "", tamanho=Pt(8))
    doc.add_paragraph()


def codigo(doc, linhas):
    p = doc.add_paragraph(style="Source")
    for i, l in enumerate(linhas):
        run = p.add_run(l)
        run.font.name = "Consolas"
        if i < len(linhas) - 1:
            run.add_break(WD_BREAK.LINE)


def converter(md_texto, doc):
    linhas = md_texto.splitlines()
    i, titulo_feito = 0, False
    while i < len(linhas):
        l = linhas[i]
        s = l.strip()
        if s.startswith("```"):
            bloco = []
            i += 1
            while i < len(linhas) and not linhas[i].strip().startswith("```"):
                bloco.append(linhas[i])
                i += 1
            codigo(doc, bloco)
            i += 1
            continue
        if s.startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                bloco.append(linhas[i])
                i += 1
            tabela(doc, bloco)
            continue
        if s.startswith(">"):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith(">"):
                bloco.append(linhas[i].strip()[1:].strip())
                i += 1
            texto, prefixo = [], None
            for b in bloco:
                m = re.fullmatch(r"\[!(\w+)\]", b)
                if m:
                    prefixo = ALERTAS.get(m.group(1).upper(), m.group(1))
                elif not b:
                    texto.append("\n")
                elif re.match(r"^(?:[-*]|\d+\.)\s", b) or b.startswith("**") and texto and texto[-1] != "\n":
                    texto.extend(["\n", b])   # item de lista/linha de rótulo dentro da citação = parágrafo próprio
                else:
                    texto.append(b)
            paragrafos = " ".join(texto).split("\n")
            for k, ptxt in enumerate(x.strip() for x in paragrafos):
                if not ptxt:
                    continue
                p = doc.add_paragraph(style="Small Text")
                if prefixo and k == 0:
                    texto_em_linha(p, "**%s.** " % prefixo)
                if ptxt.startswith("- "):
                    ptxt = "•  " + ptxt[2:]
                texto_em_linha(p, ptxt)
            continue
        if not s or s == "---":
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            nivel, txt = len(m.group(1)), m.group(2)
            if nivel == 1 and not titulo_feito:
                doc.add_paragraph("DOCUMENTO OPERACIONAL  ·  ARENAMON  ·  V2.0", style="Small Text")
                doc.add_paragraph(txt, style="Title")
                titulo_feito = True
            else:
                doc.add_paragraph(re.sub(r"`([^`]+)`", r"\1", txt), style="Heading %d" % max(1, nivel - 1))   # ## = seção, ### = subseção
            i += 1
            continue
        mb = re.match(r"^(\s*)(?:[-*]|\d+\.)\s+(.*)", l)
        if mb:
            recuo = len(mb.group(1)) // 2
            num = re.match(r"^\s*(\d+)\.", l)
            txt = mb.group(2)
            txt = re.sub(r"^\[ \]\s*", "☐ ", txt)
            # junta linhas de continuação (recuadas, sem marcador)
            i += 1
            while i < len(linhas) and linhas[i].startswith("  ") and linhas[i].strip() and not re.match(r"^\s*(?:[-*]|\d+\.)\s+", linhas[i]):
                txt += " " + linhas[i].strip()
                i += 1
            p = doc.add_paragraph(style="Bullet Compact")
            p.paragraph_format.left_indent = Pt(14 * (recuo + 1))
            texto_em_linha(p, ("%s.  " % num.group(1) if num else "•  ") + txt)
            continue
        # parágrafo comum (junta linhas consecutivas)
        txt = s
        i += 1
        while i < len(linhas) and linhas[i].strip() and not re.match(r"^\s*(#|\||>|```|[-*]\s|\d+\.\s)", linhas[i]):
            txt += " " + linhas[i].strip()
            i += 1
        texto_em_linha(doc.add_paragraph(), txt)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--md", default=MD)
    ap.add_argument("--modelo", default=MODELO)
    ap.add_argument("--saida", default=SAIDA)
    a = ap.parse_args()
    if os.path.abspath(a.saida) == os.path.abspath(a.modelo):
        sys.exit("A saída não pode ser o modelo V1 (o original é preservado).")
    doc = docx.Document(a.modelo)
    limpar_corpo(doc)
    with open(a.md, encoding="utf-8") as f:
        converter(f.read(), doc)
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        paragraph = doc.paragraphs[-1]._element
        paragraph.getparent().remove(paragraph)
    doc.core_properties.title = "ARENAMON — Manual Blender para agentes de IA (V2.0)"
    doc.core_properties.revision = 2
    doc.save(a.saida)
    print("DOCX gerado:", os.path.relpath(a.saida, RAIZ))


if __name__ == "__main__":
    main()
